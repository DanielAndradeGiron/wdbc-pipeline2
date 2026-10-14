"""Run directory layout and artifact persistence (tables, JSON, figures, checksums)."""

from __future__ import annotations

import hashlib
import json
import platform
import sys
from dataclasses import dataclass
from importlib import metadata
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

from .config import Config, config_hash
from .logging_utils import get_logger

log = get_logger("artifacts")

TRACKED_PACKAGES = (
    "numpy",
    "pandas",
    "scipy",
    "scikit-learn",
    "matplotlib",
    "seaborn",
    "joblib",
    "xgboost",
    "lightgbm",
    "catboost",
    "shap",
    "lime",
    "openpyxl",
    "python-docx",
    "PyYAML",
)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path, chunk: int = 1 << 20) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(chunk), b""):
            h.update(block)
    return h.hexdigest()


def _json_default(obj: Any) -> Any:
    if isinstance(obj, (np.integer,)):
        return int(obj)
    if isinstance(obj, (np.floating,)):
        return float(obj)
    if isinstance(obj, np.ndarray):
        return obj.tolist()
    if isinstance(obj, (np.bool_,)):
        return bool(obj)
    if isinstance(obj, Path):
        return str(obj)
    return str(obj)


@dataclass
class RunContext:
    """All paths for one run. Every stage reads/writes through this object."""

    cfg: Config
    run_dir: Path

    @classmethod
    def create(cls, cfg: Config, run_dir: Path | None = None) -> RunContext:
        if run_dir is None:
            run_dir = Path(cfg.project.output_dir) / f"{cfg.project.run_name}_{config_hash(cfg)}"
        ctx = cls(cfg=cfg, run_dir=Path(run_dir))
        for p in (
            ctx.tables,
            ctx.figures,
            ctx.models,
            ctx.metadata,
            ctx.data,
            ctx.checkpoints,
            ctx.logs,
        ):
            p.mkdir(parents=True, exist_ok=True)
        ctx._bind_config()
        return ctx

    # -- layout ---------------------------------------------------------------
    @property
    def tables(self) -> Path:
        return self.run_dir / "tables"

    @property
    def figures(self) -> Path:
        return self.run_dir / "figures"

    @property
    def models(self) -> Path:
        return self.run_dir / "models"

    @property
    def metadata(self) -> Path:
        return self.run_dir / "metadata"

    @property
    def data(self) -> Path:
        return self.run_dir / "data"

    @property
    def checkpoints(self) -> Path:
        return self.run_dir / "checkpoints"

    @property
    def logs(self) -> Path:
        return self.run_dir / "logs"

    # -- config binding -------------------------------------------------------
    def _bind_config(self) -> None:
        """Refuse to mix results produced under different analysis plans."""
        h = config_hash(self.cfg)
        stamp = self.metadata / "config_hash.txt"
        if stamp.exists() and stamp.read_text().strip() != h:
            raise RuntimeError(
                f"Run directory {self.run_dir} was created with config hash "
                f"{stamp.read_text().strip()}, current config hash is {h}. "
                "Use a new run directory (or the original configuration)."
            )
        stamp.write_text(h)
        self.save_json(self.cfg.to_dict(), "config_resolved")
        self.save_json(runtime_versions(), "versions_runtime")

    # -- persistence ----------------------------------------------------------
    def save_json(self, obj: Any, name: str) -> Path:
        path = self.metadata / f"{name}.json"
        path.write_text(json.dumps(obj, indent=2, default=_json_default), encoding="utf-8")
        return path

    def load_json(self, name: str) -> Any:
        path = self.metadata / f"{name}.json"
        if not path.exists():
            raise FileNotFoundError(f"{path} not found — run the stage that produces it first.")
        return json.loads(path.read_text(encoding="utf-8"))

    def has_json(self, name: str) -> bool:
        return (self.metadata / f"{name}.json").exists()

    def save_table(self, df: pd.DataFrame, name: str) -> Path:
        csv = self.tables / f"{name}.csv"
        df.to_csv(csv, index=False, encoding="utf-8-sig")
        if self.cfg.runtime.excel_export:
            _write_excel(df, self.tables / f"{name}.xlsx")
        return csv

    def load_table(self, name: str) -> pd.DataFrame:
        path = self.tables / f"{name}.csv"
        if not path.exists():
            raise FileNotFoundError(f"{path} not found — run the stage that produces it first.")
        return pd.read_csv(path, encoding="utf-8-sig")

    def save_figure(self, fig: Any, name: str) -> list[Path]:
        import matplotlib.pyplot as plt

        out = []
        for ext in self.cfg.runtime.figure_formats:
            path = self.figures / f"{name}.{ext}"
            fig.savefig(path, bbox_inches="tight", facecolor="white", dpi=self.cfg.runtime.dpi)
            out.append(path)
        plt.close(fig)
        return out


def _write_excel(df: pd.DataFrame, path: Path, max_rows: int = 50_000) -> None:
    """Formatted, filterable Excel export.

    Column widths are computed from pandas (vectorised) instead of iterating
    over every cell; very large tables are kept as CSV only.
    """
    if len(df) > max_rows:
        log.debug("Skipping Excel export for %s (%d rows > %d)", path.name, len(df), max_rows)
        return
    try:
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="data")
            ws = writer.sheets["data"]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            fill = PatternFill("solid", fgColor="17365D")
            for cell in ws[1]:
                cell.fill = fill
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for i, col in enumerate(df.columns, start=1):
                body = df[col].astype(str).str.len().max() if len(df) else 0
                width = min(34, max(10, int(max(len(str(col)), body or 0)) + 2))
                ws.column_dimensions[get_column_letter(i)].width = width
    except Exception as exc:  # Excel is a convenience copy; CSV is canonical.
        log.warning("Excel export failed for %s: %r", path.name, exc)


def runtime_versions() -> dict[str, str]:
    versions = {"python": sys.version.split()[0], "platform": platform.platform()}
    for pkg in TRACKED_PACKAGES:
        try:
            versions[pkg] = metadata.version(pkg)
        except metadata.PackageNotFoundError:
            versions[pkg] = "not installed"
    return versions


def checksum_manifest(root: Path, exclude: tuple[str, ...] = ("checkpoints",)) -> pd.DataFrame:
    rows = []
    for fp in sorted(p for p in root.rglob("*") if p.is_file()):
        rel = fp.relative_to(root)
        if rel.parts and rel.parts[0] in exclude:
            continue
        rows.append(
            {"relative_path": rel.as_posix(), "bytes": fp.stat().st_size, "sha256": sha256_file(fp)}
        )
    return pd.DataFrame(rows)
