"""Quality gates, Markdown/Word reports, checksum manifest and result package."""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

from .artifacts import RunContext, checksum_manifest
from .data import HoldoutGuard, load_development
from .logging_utils import get_logger
from .models import audit_stacking_architecture, build_model

log = get_logger("report")


class QualityGateError(AssertionError):
    pass


def quality_gates(ctx: RunContext) -> dict:
    cfg = ctx.cfg
    split = ctx.load_json("split_manifest")
    dev = load_development(ctx)
    hold = HoldoutGuard(ctx).read_after_opening()
    fold_df = ctx.load_table("nested_repeated_cv_by_fold")
    oof = ctx.load_table("nested_repeated_cv_predictions")
    th_oof = ctx.load_table("threshold_oof_predictions")
    preds = ctx.load_table("final_internal_holdout_predictions")
    ci = ctx.load_table("table_06_holdout_bootstrap_ci")
    exact = ctx.load_table("table_07_holdout_exact_binomial_ci")
    usable = ctx.load_json("benchmark_models")
    frozen = ctx.load_json("frozen_configuration_before_holdout")
    selection = ctx.load_json("selection_rule")
    access = (
        (ctx.metadata / "holdout_access_log.jsonl").read_text(encoding="utf-8").strip().splitlines()
    )

    checks: dict[str, object] = {
        "split_disjoint": set(dev.ids).isdisjoint(hold.ids),
        "split_complete": len(dev) + len(hold) == split["n_total"],
        "benchmark_probabilities_valid": bool(oof["probability"].between(0, 1).all()),
        "all_usable_models_benchmarked": set(fold_df["model"]) == set(usable["usable"]),
        "benchmark_candidates_calibrated": bool(fold_df["calibrated"].all()),
        "threshold_oof_one_per_development_record": bool(
            th_oof["source_row_id"].is_unique and len(th_oof) == len(dev)
        ),
        "threshold_oof_probabilities_valid": bool(th_oof["probability_oof"].between(0, 1).all()),
        "holdout_predictions_one_per_record": bool(
            preds["source_row_id"].is_unique and len(preds) == len(hold)
        ),
        "selected_final_algorithm_consistent": frozen["final_model"] == selection["selected_model"],
        "both_threshold_ci_available": set(ci["evaluation"]) == {"primary", "secondary_OOF"},
        "exact_binomial_ci_available": {"sensitivity", "specificity"} <= set(exact["metric"]),
    }
    if "Stacking" in usable["usable"]:
        checks["stacking_architecture_passed"] = audit_stacking_architecture(
            build_model("Stacking", dev.X, cfg.project.seed, cfg), cfg
        )
    failed = [k for k, v in checks.items() if not bool(v)]
    checks["informational"] = {
        "validation_scope": "locked internal holdout only",
        "models_skipped_missing_dependency": usable["skipped"],
        "holdout_openings": len(access),
        "holdout_single_opening": len(access) == 1,
    }
    checks["all_required_gates_passed"] = not failed
    ctx.save_json(checks, "quality_checks")
    if failed:
        raise QualityGateError(f"Quality gates failed: {failed}")
    if len(access) > 1:
        log.warning("The holdout was opened %d times — disclose this in any report.", len(access))
    return checks


def _fmt(x: float) -> str:
    return f"{x:.4f}"


def markdown_report(ctx: RunContext) -> str:
    cfg = ctx.cfg
    split = ctx.load_json("split_manifest")
    sel = ctx.load_json("selection_rule")
    st = ctx.load_json("stats_summary")
    frozen = ctx.load_json("frozen_configuration_before_holdout")
    expl = ctx.load_json("explainability_status") if ctx.has_json("explainability_status") else {}
    usable = ctx.load_json("benchmark_models")
    summary = ctx.load_table("table_01_repeated_nested_cv_model_summary")
    hold = ctx.load_table("table_05_final_internal_holdout_metrics")
    pri, sec = hold.iloc[0], hold.iloc[1]

    lines = [
        "# Automated internal-evaluation technical report",
        "",
        f"- Run directory: `{ctx.run_dir.as_posix()}`  ",
        f"- Config hash: `{frozen['config_hash']}` — data source: `{cfg.data.source}`",
        "",
        "## Analysis design",
        f"- Analytical cohort **{split['n_total']}**; development **{split['n_development']}**; "
        f"locked internal holdout **{split['n_holdout']}** (seed {split['seed']}).",
        f"- Benchmark: **{cfg.cv.outer_splits}-fold × {cfg.cv.outer_repeats}-repeat nested CV**, inner "
        f"{cfg.cv.inner_splits}-fold RandomizedSearchCV (n_iter={cfg.cv.search_n_iter}, {cfg.cv.scoring}), "
        f"fold-local {cfg.cv.calibration_method} calibration for every candidate.",
        f"- Candidates: {', '.join(usable['usable'])}."
        + (
            f" Skipped (missing dependency): {', '.join(usable['skipped'])}."
            if usable["skipped"]
            else ""
        ),
        f"- Selected and final algorithm: **{frozen['final_model']}** (practical-tolerance cascade, "
        "computational tie-break).",
    ]
    if sel.get("note"):
        lines.append(f"- Note: {sel['note']}")
    lines += [
        f"- Friedman test (exploratory): χ² = **{st['friedman_chi_square']:.3f}**, p = **{st['friedman_p']:.4g}**; "
        f"Holm-significant corrected paired comparisons: **{st['n_significant_holm']}**.",
        "",
        "## Repeated nested CV summary",
        "",
        "| Rank | Model | Weighted F1 | MCC | ROC-AUC | Brier | ECE | Fit (s) |",
        "|---:|---|---|---|---|---|---|---:|",
    ]
    for _, r in summary.iterrows():
        mark = " **(selected)**" if r["selected"] else ""
        lines.append(
            f"| {r['descriptive_rank']} | {r['model']}{mark} | {_fmt(r['f1_weighted_mean'])} ± "
            f"{_fmt(r['f1_weighted_sd'])} | {_fmt(r['mcc_mean'])} | {_fmt(r['roc_auc_mean'])} | "
            f"{_fmt(r['brier_mean'])} | {_fmt(r['ece_mean'])} | {r['fit_time_mean']:.2f} |"
        )
    lines += [
        "",
        "## Final locked internal-holdout evaluation",
        f"- Final hyperparameters: `{json.dumps(frozen['best_params'], default=str)}`",
        f"- Primary threshold **{pri['threshold']:.2f}**: accuracy {_fmt(pri['accuracy'])}, sensitivity "
        f"{_fmt(pri['sensitivity'])}, specificity {_fmt(pri['specificity'])}, weighted F1 "
        f"{_fmt(pri['f1_weighted'])}, MCC {_fmt(pri['mcc'])}.",
        f"- Secondary fully nested OOF threshold **{sec['threshold']:.2f}** ({frozen['threshold_rule']}): "
        f"accuracy {_fmt(sec['accuracy'])}, sensitivity {_fmt(sec['sensitivity'])}, specificity "
        f"{_fmt(sec['specificity'])}, weighted F1 {_fmt(sec['f1_weighted'])}, MCC {_fmt(sec['mcc'])}.",
        f"- Threshold-independent: ROC-AUC {_fmt(pri['roc_auc'])}, average precision "
        f"{_fmt(pri['average_precision'])}, Brier {_fmt(pri['brier'])}, ECE {_fmt(pri['ece'])}.",
        "",
        "## Interpretation constraints",
        "- The holdout is internal; results do not demonstrate external transportability, prospective "
        "clinical utility or causal effects.",
        "- The secondary threshold was derived on development OOF predictions and transferred to a model "
        "refitted on all development data; it is provisional until externally validated.",
        "- DCA is exploratory and must be reassessed under setting-specific prevalence and utilities.",
        f"- Explainability is post hoc. LIME is reportable only when local R² ≥ {cfg.explain.lime_min_fidelity:.2f}. "
        f"Status: `{json.dumps(expl)}`.",
    ]
    if cfg.cv.outer_repeats == 1 and cfg.cv.outer_splits <= 3:
        lines.insert(1, "> ⚠️ **QUICK MODE — structural test only. Do not report these numbers.**")
    return "\n".join(lines) + "\n"


def word_report(ctx: RunContext, path: Path) -> None:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt
    except ImportError:
        log.warning("python-docx not installed; Word report skipped.")
        return
    split = ctx.load_json("split_manifest")
    frozen = ctx.load_json("frozen_configuration_before_holdout")
    st = ctx.load_json("stats_summary")
    summary = ctx.load_table("table_01_repeated_nested_cv_model_summary")
    hold = ctx.load_table("table_05_final_internal_holdout_metrics")
    cfg = ctx.cfg

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    doc.add_heading(
        "WDBC ensemble benchmarking — technical results", 0
    ).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "Leakage-resistant repeated nested validation with a locked internal holdout"
    ).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("Analysis design", 1)
    doc.add_paragraph(
        f"The analytical cohort contained {split['n_total']} observations: {split['n_development']} for "
        f"development and {split['n_holdout']} in the locked internal holdout. Benchmarking used "
        f"{cfg.cv.outer_splits}-fold cross-validation repeated {cfg.cv.outer_repeats} times with "
        "hyperparameter tuning and probability calibration repeated inside every outer training partition."
    )
    doc.add_heading("Repeated nested cross-validation", 1)
    doc.add_paragraph(
        f"{frozen['final_model']} was selected by the prespecified practical-tolerance hierarchy and "
        f"computational tie-break. The Friedman test was exploratory (χ²={st['friedman_chi_square']:.3f}, "
        f"p={st['friedman_p']:.4g}) and does not establish superiority or equivalence."
    )

    def table(header: list[str], rows: list[list[str]]) -> None:
        t = doc.add_table(rows=1, cols=len(header))
        t.style = "Light Shading Accent 1"
        for c, v in zip(t.rows[0].cells, header):
            c.text = v
        for r in rows:
            for c, v in zip(t.add_row().cells, r):
                c.text = str(v)

    table(
        ["Model", "Weighted F1", "MCC", "ROC-AUC", "AP", "Brier", "Time (s)"],
        [
            [
                r["model"],
                f"{r['f1_weighted_mean']:.4f} ± {r['f1_weighted_sd']:.4f}",
                _fmt(r["mcc_mean"]),
                _fmt(r["roc_auc_mean"]),
                _fmt(r["average_precision_mean"]),
                _fmt(r["brier_mean"]),
                f"{r['fit_time_mean']:.2f}",
            ]
            for _, r in summary.iterrows()
        ],
    )
    figs = ctx.figures
    if (figs / "figure_01_nested_cv_weighted_f1.png").exists():
        doc.add_picture(str(figs / "figure_01_nested_cv_weighted_f1.png"), width=Inches(6.4))
    doc.add_heading("Locked internal-holdout evaluation", 1)
    table(
        ["Evaluation", "Threshold", "Accuracy", "Sensitivity", "Specificity", "Weighted F1", "MCC"],
        [
            [
                r["evaluation"],
                f"{r['threshold']:.2f}",
                _fmt(r["accuracy"]),
                _fmt(r["sensitivity"]),
                _fmt(r["specificity"]),
                _fmt(r["f1_weighted"]),
                _fmt(r["mcc"]),
            ]
            for _, r in hold.iterrows()
        ],
    )
    for name in (
        "figure_03_holdout_confusion_matrices",
        "figure_04_holdout_roc_pr",
        "figure_05_holdout_calibration",
        "figure_06_holdout_dca",
        "figure_07a_permutation_importance",
        "figure_07_shap_global_summary",
    ):
        fp = figs / f"{name}.png"
        if fp.exists():
            doc.add_picture(str(fp), width=Inches(6.2))
    doc.add_heading("Interpretation limits", 1)
    doc.add_paragraph(
        "These findings establish internal performance only. They do not demonstrate external "
        "transportability, prospective clinical utility, or causal effects. Decision-curve and "
        "explainability analyses are post hoc; LIME is suppressed when its local surrogate fidelity is "
        "below the prespecified threshold."
    )
    doc.save(str(path))


def run_report(ctx: RunContext, make_zip: bool = True) -> Path:
    quality_gates(ctx)
    md = markdown_report(ctx)
    (ctx.run_dir / "TECHNICAL_RESULTS_REPORT.md").write_text(md, encoding="utf-8")
    word_report(ctx, ctx.run_dir / "PUBLICATION_RESULTS_REPORT.docx")
    (ctx.run_dir / "README.md").write_text(
        "# Output package\n\n"
        "- `tables/`: canonical CSV tables (+ formatted Excel copies).\n"
        "- `figures/`: publication figures.\n"
        "- `models/`: frozen calibrated model (`joblib`).\n"
        "- `metadata/`: resolved config + hash, versions, data/split manifests, frozen decisions, "
        "holdout access log, quality gates, explainability status.\n"
        "- `data/`: dataset snapshot, development set and locked holdout (with SHA-256 in the split manifest).\n"
        "- `logs/`: full pipeline log.\n"
        "- `TECHNICAL_RESULTS_REPORT.md`, `PUBLICATION_RESULTS_REPORT.docx`: automated reports.\n"
        "- `artifact_checksum_manifest.csv`: SHA-256 of every file above.\n",
        encoding="utf-8",
    )
    manifest = checksum_manifest(ctx.run_dir)
    manifest = manifest[manifest["relative_path"] != "artifact_checksum_manifest.csv"]
    manifest.to_csv(ctx.run_dir / "artifact_checksum_manifest.csv", index=False)
    log.info("Report written: %s", ctx.run_dir / "TECHNICAL_RESULTS_REPORT.md")
    if make_zip:
        archive = ctx.run_dir.with_suffix(".zip")
        with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as zf:
            for fp in sorted(p for p in ctx.run_dir.rglob("*") if p.is_file()):
                rel = fp.relative_to(ctx.run_dir)
                if rel.parts[0] != "checkpoints":
                    zf.write(fp, Path(ctx.run_dir.name) / rel)
        log.info("Result package: %s", archive)
        return archive
    return ctx.run_dir
